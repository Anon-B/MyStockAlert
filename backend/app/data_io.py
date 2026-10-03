from __future__ import annotations
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import PortfolioHolding, PortfolioTransaction
from .finance import FeeBreakdown, calculate_transaction, convert_to_thb


def _style(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
    for col in ws.columns:
        width = min(max(max(len(str(c.value or "")) for c in col) + 2, 10), 32)
        ws.column_dimensions[col[0].column_letter].width = width


def _book(sheet: str, headers: list[str], rows: list[list]):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(row)
    _style(ws)
    return wb


def portfolio_workbook(rows, template=False):
    headers = ["symbol", "market", "name", "quantity", "avg_cost", "currency"]
    if not template:
        headers += ["current_price", "market_value", "unrealized_pnl"]
        data = [[x["symbol"], x["market"], x.get("name", ""), float(x["quantity"]), float(x["avg_cost"]), x["currency"], x.get("current_price"), x.get("market_value"), x.get("unrealized_pnl")] for x in rows]
    else:
        headers = ["symbol", "market", "quantity", "avg_cost", "currency"]
        data = []
    wb = _book("Portfolio", headers, data)
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Required", "Description"])
    for x in [("symbol", "Yes", "Ticker symbol, e.g. PTT or AAPL"), ("market", "Yes", "TH or US"), ("quantity", "Yes", "Quantity > 0"), ("avg_cost", "Yes", "Average cost per share > 0"), ("currency", "Yes", "THB or USD")]: ins.append(x)
    _style(ins)
    ex = wb.create_sheet("Examples")
    ex.append(headers); ex.append(["PTT", "TH", 100, 35.5, "THB"]); ex.append(["AAPL", "US", 10, 280, "USD"]); _style(ex)
    return wb


def transaction_workbook(rows, template=False):
    headers = ["external_id", "transaction_date", "symbol", "market", "side", "quantity", "price", "currency", "fee", "tax", "note"]
    data = []
    if not template:
        for item in rows:
            x, holding = item
            fee = sum(Decimal(str(v or 0)) for v in [x.commission, x.trading_fee, x.clearing_fee, x.regulatory_fee, x.cat_fee, x.sec_fee, x.taf_fee])
            data.append([x.import_external_id or "", x.executed_at.replace(tzinfo=None) if x.executed_at.tzinfo else x.executed_at, holding.symbol, holding.market, x.side, float(x.quantity), float(x.execution_price), x.currency, float(fee), float(x.vat or 0), ""])
    wb = _book("Transactions", headers, data)
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Required", "Description"])
    for x in [("external_id", "Yes", "Unique import ID; re-importing the same ID is skipped"), ("transaction_date", "Yes", "Trade date/time"), ("symbol", "Yes", "Ticker symbol"), ("market", "Yes", "TH or US"), ("side", "Yes", "BUY or SELL"), ("quantity", "Yes", "Quantity > 0"), ("price", "Yes", "Execution price > 0"), ("currency", "Yes", "THB or USD"), ("fee", "No", "Total fee; imported as commission"), ("tax", "No", "Tax/VAT amount"), ("note", "No", "Reference note")]: ins.append(x)
    _style(ins)
    ex = wb.create_sheet("Examples")
    ex.append(headers); ex.append(["TX-000001", datetime(2026, 10, 1, 10, 0), "PTT", "TH", "BUY", 100, 35.5, "THB", 5, 0, "ตัวอย่างซื้อ"]); ex.append(["TX-000002", datetime(2026, 10, 2, 10, 0), "AAPL", "US", "BUY", 10, 280, "USD", 2, 0, "ตัวอย่างซื้อ"]); _style(ex)
    return wb


def as_bytes(wb):
    out = BytesIO(); wb.save(out); return out.getvalue()


def _read_rows(content: bytes, sheet: str):
    wb = load_workbook(BytesIO(content), data_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f"missing sheet: {sheet}")
    ws = wb[sheet]
    values = list(ws.iter_rows(values_only=True))
    if not values: return []
    headers = [str(x).strip().lower() if x is not None else "" for x in values[0]]
    return [{headers[i]: row[i] if i < len(row) else None for i in range(len(headers))} for row in values[1:] if any(v is not None and str(v).strip() != "" for v in row)]


def _dec(v, name, errors, row, gt_zero=False):
    if v is None or str(v).strip() == "": errors.append(f"{name} is required"); return None
    try: d = Decimal(str(v).strip())
    except (InvalidOperation, ValueError): errors.append(f"{name} must be numeric"); return None
    if gt_zero and d <= 0: errors.append(f"{name} must be greater than 0")
    if not gt_zero and d < 0: errors.append(f"{name} must not be negative")
    return d


def preview_portfolio(db: Session, user_id, content: bytes):
    try: rows = _read_rows(content, "Portfolio")
    except Exception as e: return {"valid": 0, "errors": [{"row": 0, "message": str(e)}], "warnings": []}
    errors=[]; warnings=[]; seen=set(); valid=0
    for n,r in enumerate(rows,2):
        e=[]; symbol=str(r.get("symbol") or "").strip().upper(); market=str(r.get("market") or "").strip().upper(); currency=str(r.get("currency") or "").strip().upper()
        q=_dec(r.get("quantity"),"quantity",e,n,True); ac=_dec(r.get("avg_cost"),"avg_cost",e,n,True)
        if market not in {"TH","US"}: e.append("market must be TH or US")
        if currency not in {"THB","USD"}: e.append("currency must be THB or USD")
        key=(market,symbol)
        if not symbol: e.append("symbol is required")
        if key in seen: e.append("duplicate symbol/market in file")
        seen.add(key)
        existing=db.scalar(select(PortfolioHolding).where(PortfolioHolding.user_id==user_id,PortfolioHolding.market==market,PortfolioHolding.symbol==symbol)) if symbol and market in {"TH","US"} else None
        if existing: warnings.append({"row":n,"message":f"{market}:{symbol} already exists and will be updated"})
        if not e: valid+=1
        errors.extend({"row":n,"message":x} for x in e)
    return {"valid":valid,"total":len(rows),"errors":errors,"warnings":warnings}


def preview_transactions(db: Session, user_id, content: bytes):
    try: rows = _read_rows(content, "Transactions")
    except Exception as e: return {"valid": 0, "errors": [{"row": 0, "message": str(e)}], "warnings": []}
    errors=[]; warnings=[]; seen=set(); valid=0
    holdings={ (x.market,x.symbol): x for x in db.scalars(select(PortfolioHolding).where(PortfolioHolding.user_id==user_id)).all() }
    for n,r in enumerate(rows,2):
        e=[]; ext=str(r.get("external_id") or "").strip(); symbol=str(r.get("symbol") or "").strip().upper(); market=str(r.get("market") or "").strip().upper(); side=str(r.get("side") or "").strip().upper(); currency=str(r.get("currency") or "").strip().upper()
        if not ext: e.append("external_id is required")
        if ext in seen: e.append("duplicate external_id in file")
        seen.add(ext)
        if not symbol: e.append("symbol is required")
        if market not in {"TH","US"}: e.append("market must be TH or US")
        if side not in {"BUY","SELL"}: e.append("side must be BUY or SELL")
        if currency not in {"THB","USD"}: e.append("currency must be THB or USD")
        q=_dec(r.get("quantity"),"quantity",e,n,True); p=_dec(r.get("price"),"price",e,n,True)
        dt=r.get("transaction_date")
        if not isinstance(dt, datetime):
            try: dt=datetime.fromisoformat(str(dt).replace("Z","+00:00"))
            except Exception: e.append("transaction_date must be a valid date/time")
        if ext and db.scalar(select(PortfolioTransaction).where(PortfolioTransaction.user_id==user_id,PortfolioTransaction.import_external_id==ext)): warnings.append({"row":n,"message":"external_id already imported; row will be skipped"})
        if symbol and market in {"TH","US"} and (market,symbol) not in holdings: warnings.append({"row":n,"message":f"{market}:{symbol} is not in Portfolio; it will be created from this transaction"})
        if not e: valid+=1
        errors.extend({"row":n,"message":x} for x in e)
    return {"valid":valid,"total":len(rows),"errors":errors,"warnings":warnings}


def import_portfolio(db: Session, user_id, content: bytes):
    p=preview_portfolio(db,user_id,content)
    if p["errors"]: return {**p,"imported":0,"skipped":0,"failed":len(p["errors"])}
    rows=_read_rows(content,"Portfolio"); imported=0
    for r in rows:
        market=str(r["market"]).strip().upper(); symbol=str(r["symbol"]).strip().upper(); currency=str(r["currency"]).strip().upper()
        row=db.scalar(select(PortfolioHolding).where(PortfolioHolding.user_id==user_id,PortfolioHolding.market==market,PortfolioHolding.symbol==symbol))
        if row: row.quantity=Decimal(str(r["quantity"])); row.average_cost=Decimal(str(r["avg_cost"])); row.currency=currency; row.enabled=True
        else: db.add(PortfolioHolding(user_id=user_id,market=market,symbol=symbol,quantity=Decimal(str(r["quantity"])),average_cost=Decimal(str(r["avg_cost"])),currency=currency,enabled=True))
        imported+=1
    db.commit(); return {**p,"imported":imported,"skipped":0,"failed":0}


def import_transactions(db: Session, user_id, content: bytes):
    p=preview_transactions(db,user_id,content)
    if p["errors"]: return {**p,"imported":0,"skipped":0,"failed":len(p["errors"])}
    rows=_read_rows(content,"Transactions"); imported=skipped=0; holdings={ (x.market,x.symbol): x for x in db.scalars(select(PortfolioHolding).where(PortfolioHolding.user_id==user_id)).all() }
    for r in rows:
        ext=str(r["external_id"]).strip()
        if db.scalar(select(PortfolioTransaction).where(PortfolioTransaction.user_id==user_id,PortfolioTransaction.import_external_id==ext)): skipped+=1; continue
        market=str(r["market"]).strip().upper(); symbol=str(r["symbol"]).strip().upper(); currency=str(r["currency"]).strip().upper(); side=str(r["side"]).strip().upper(); q=Decimal(str(r["quantity"])); price=Decimal(str(r["price"]))
        holding=holdings.get((market,symbol))
        if not holding:
            holding=PortfolioHolding(user_id=user_id,market=market,symbol=symbol,quantity=Decimal("0"),average_cost=Decimal("0"),currency=currency,enabled=True); db.add(holding); db.flush(); holdings[(market,symbol)]=holding
        fee=Decimal(str(r.get("fee") or 0)); tax=Decimal(str(r.get("tax") or 0)); fees=FeeBreakdown(commission=fee,vat=tax)
        trading_value,total_fees,net_amount=calculate_transaction(side,q,price,fees)
        if side=="SELL":
            if q>holding.quantity: raise ValueError(f"{market}:{symbol} sell quantity exceeds holding quantity")
            holding.quantity-=q
        else:
            old=holding.quantity*holding.average_cost; holding.quantity+=q; holding.average_cost=(old+net_amount)/holding.quantity
        dt=r.get("transaction_date")
        if not isinstance(dt,datetime): dt=datetime.fromisoformat(str(dt).replace("Z","+00:00"))
        if dt.tzinfo is None: dt=dt.replace(tzinfo=timezone.utc)
        thb=trading_value if currency=="THB" else None
        net_thb=net_amount if currency=="THB" else None
        db.add(PortfolioTransaction(user_id=user_id,holding_id=holding.id,import_external_id=ext,idempotency_key=f"import:{ext}",side=side,status="FILLED",order_id=None,quantity=q,execution_price=price,trading_value=trading_value,trading_value_thb=thb,commission=fee,trading_fee=0,clearing_fee=0,regulatory_fee=0,cat_fee=0,sec_fee=0,taf_fee=0,vat=tax,fx_rate=None,net_amount=net_amount,net_amount_thb=net_thb,currency=currency,executed_at=dt)); imported+=1
    db.commit(); return {**p,"imported":imported,"skipped":skipped,"failed":0}
